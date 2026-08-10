from openpyxl import Workbook
import pytest

from scripts.subway_translation import load_kric_station_translations


def test_loads_all_supported_languages_and_normalizes_station_suffix(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "역명(한글)",
        "역명(영문)",
        "역명(일본어)",
        "역명(중국어 간체)",
        "역명(중국어 번체) ",
    ])
    sheet.append(["김포공항역", "Gimpo Int'l Airport", "キンポゴンハン", "金浦机场", "金浦空港"])
    path = tmp_path / "seohae.xlsx"
    workbook.save(path)
    workbook.close()

    result = load_kric_station_translations(path, [("S13", "김포공항")])

    assert {(item["language"], item["name"]) for item in result} == {
        ("ko", "김포공항"),
        ("en", "Gimpo Int'l Airport"),
        ("ja", "キンポゴンハン"),
        ("zh-Hans", "金浦机场"),
        ("zh-Hant", "金浦空港"),
    }
    assert all(item["station_id"] == "S13" for item in result)
    assert all(item["source"] == "KRIC_20250630" for item in result)


def test_rejects_stations_missing_from_official_workbook(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["역명", "영어명", "일본어", "중국어간체", "중국어번체"])
    sheet.append(["안산", "Ansan", "アンサン", "安山", "安山"])
    path = tmp_path / "line4.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match=r"K449 \(한대앞\)"):
        load_kric_station_translations(path, [("K449", "한대앞")])

    assert load_kric_station_translations(
        path,
        [("K449", "한대앞")],
        strict=False,
    ) == []


def test_maps_known_kric_aliases_without_replacing_korean_name(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["역명", "영어명", "일본어", "중국어간체", "중국어번체"])
    sheet.append([
        "총신대입구(이수)",
        "Chongshin Univ. (Isu)",
        "チョンシンデイック",
        "总神大入口",
        "總神大入口",
    ])
    path = tmp_path / "line4.xlsx"
    workbook.save(path)
    workbook.close()

    result = load_kric_station_translations(path, [("K432", "이수")])

    assert next(item["name"] for item in result if item["language"] == "ko") == "이수"
    assert next(item["name"] for item in result if item["language"] == "en") == "Chongshin Univ. (Isu)"


def test_supplements_missing_traditional_chinese_and_renamed_station(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["역명", "영어명", "일본어", "중국어간체", "중국어번체"])
    sheet.append(["당고개", "Danggogae", "タンゴゲ", "堂岭", None])
    sheet.append(["선바위", "Seonbawi", "ソンバウィ", "立岩", None])
    path = tmp_path / "line4.xlsx"
    workbook.save(path)
    workbook.close()

    result = load_kric_station_translations(
        path,
        [("K409", "불암산"), ("K435", "선바위")],
    )

    buramsan = {item["language"]: item for item in result if item["station_id"] == "K409"}
    assert buramsan["en"]["name"] == "Buramsan"
    assert buramsan["ja"]["name"] == "プラムサン"
    assert buramsan["zh-Hant"]["name"] == "佛巖山"
    assert buramsan["en"]["source"] == "MANUAL_REVIEW_20260810"

    seonbawi = {item["language"]: item for item in result if item["station_id"] == "K435"}
    assert seonbawi["zh-Hant"]["name"] == "立岩"
    assert seonbawi["zh-Hant"]["source"] == "MANUAL_REVIEW_20260810"
