from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


HEADER_ALIASES = {
    "ko": ("역명", "역명(한글)"),
    "en": ("영어명", "역명(영문)"),
    "ja": ("일본어", "역명(일본어)"),
    "zh-Hans": ("중국어간체", "역명(중국어간체)"),
    "zh-Hant": ("중국어번체", "역명(중국어번체)"),
}

STATION_NAME_ALIASES = {
    "과천정부청사": "정부과천청사",
    "이수": "총신대입구(이수)",
    "동작": "동작(현충원)",
    "이촌": "이촌(국립중앙박물관)",
    "숙대입구": "숙대입구(갈월)",
    "회현": "회현(남대문시장)",
    "한성대입구": "한성대입구(삼선교)",
    "성신여대입구": "성신여대입구(돈암)",
    "수유": "수유(강북구청)",
    "불암산": "당고개",
    "숭의": "숭의(인하대병원)",
    "수원시청": "수원시청(경기도문화의전당)",
    "영통": "영통(경희대)",
    "상갈": "상갈(루터대학교)",
    "기흥": "기흥(백남준아트센터)",
    "죽전": "죽전(단국대)",
    "미금": "미금(분당서울대병원)",
    "수내": "수내(한국잡월드)",
    "이매": "이매(성남아트센터)",
    "복정": "복정(동서울대학)",
    "대모산": "대모산입구",
    "선정릉": "선정릉(한국과학창의재단)",
}

# Verified supplements for renamed stations and blank cells in the KRIC workbook.
STATION_TRANSLATION_OVERRIDES = {
    "불암산": {
        "ko": "불암산",
        "en": "Buramsan",
        "ja": "プラムサン",
        "zh-Hans": "佛岩山",
        "zh-Hant": "佛巖山",
    },
    "당고개": {"zh-Hant": "堂嶺"},
    "선바위": {"zh-Hant": "立岩"},
    "범계": {"zh-Hant": "凡溪"},
    "서울숲": {"zh-Hant": "首爾林"},
}


def load_kric_station_translations(
    workbook_path: str | Path,
    stations: Iterable[tuple[str, str]],
    source: str = "KRIC_20250630",
    strict: bool = True,
) -> list[dict]:
    """Map an official KRIC workbook to HYUabot station IDs by Korean station name."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        indexes = {
            language: _find_header_index(headers, aliases)
            for language, aliases in HEADER_ALIASES.items()
        }
        official_names: dict[str, dict[str, str]] = {}
        for row in rows:
            korean_name = _normalize_station_name(_cell(row, indexes["ko"]))
            if not korean_name:
                continue
            official_names[korean_name] = {
                language: _cell(row, index)
                for language, index in indexes.items()
                if _cell(row, index)
            }

        translations: list[dict] = []
        missing_stations: list[str] = []
        for station_id, korean_name in stations:
            normalized_name = _normalize_station_name(korean_name)
            official_name = STATION_NAME_ALIASES.get(normalized_name, normalized_name)
            names = official_names.get(_normalize_station_name(official_name))
            if names is None:
                missing_stations.append(f"{station_id} ({korean_name})")
                continue
            overrides = STATION_TRANSLATION_OVERRIDES.get(normalized_name, {})
            names.update(overrides)
            for language, name in names.items():
                if language == "ko":
                    name = korean_name
                translations.append(
                    dict(
                        station_id=station_id,
                        language=language,
                        name=name,
                        source="MANUAL_REVIEW_20260810" if language in overrides else source,
                        is_verified=True,
                    ),
                )
        if strict and missing_stations:
            missing = ", ".join(missing_stations)
            raise ValueError(f"Stations missing from KRIC workbook: {missing}")
        return translations
    finally:
        workbook.close()


def _find_header_index(headers: list[str], aliases: tuple[str, ...]) -> int:
    normalized_headers = [_normalize_header(header) for header in headers]
    for alias in aliases:
        normalized_alias = _normalize_header(alias)
        if normalized_alias in normalized_headers:
            return normalized_headers.index(normalized_alias)
    raise ValueError(f"Required KRIC column is missing: {aliases}")


def _normalize_header(header: str) -> str:
    return "".join(header.split())


def _cell(row: tuple, index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _normalize_station_name(name: str) -> str:
    normalized = name.strip()
    return normalized[:-1] if normalized.endswith("역") else normalized
